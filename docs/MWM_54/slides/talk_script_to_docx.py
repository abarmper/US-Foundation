#!/usr/bin/env python3
"""Render talk_script.md as talk_script.docx.

The .md stays the source of truth (edit it and re-run); this script only
handles the constructs that file uses: H1/H2 headings, blockquote spoken
text, one pipe table, horizontal rules, plain paragraphs, and inline
**bold** / *italic* markup. Styling matches the deck: Calibri, ink #1A1A1A,
accent #6B4E8E.

Usage:  python3 talk_script_to_docx.py
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SRC = HERE / "talk_script.md"
OUT = HERE / "talk_script.docx"

INK = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0x6B, 0x4E, 0x8E)
MUTED = RGBColor(0x66, 0x66, 0x66)

_MARKUP = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*)")


def add_runs(p, txt, size, color=INK, bold=False, italic=False):
    for part in _MARKUP.split(txt):
        if not part:
            continue
        b, it, seg = bold, italic, part
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            b, seg = True, part[2:-2]
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            it, seg = True, part[1:-1]
        r = p.add_run(seg)
        r.font.name = "Calibri"
        r.font.size = Pt(size)
        r.font.bold = b
        r.font.italic = it
        r.font.color.rgb = color


def main():
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Inches(0.8)
        sec.left_margin = sec.right_margin = Inches(1.0)

    i = 0
    while i < len(lines):
        line = lines[i]

        if line.startswith("# "):                                # title
            p = doc.add_paragraph()
            add_runs(p, line[2:], 22, bold=True)
            p.space_after = None
            p.paragraph_format.space_after = Pt(10)
            i += 1

        elif line.startswith("## "):                             # section
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            add_runs(p, line[3:], 14.5, color=ACCENT, bold=True)
            i += 1

        elif line.startswith(">"):                               # blockquote
            block, i = [], i
            while i < len(lines) and lines[i].startswith(">"):
                block.append(lines[i].lstrip("> ").rstrip())
                i += 1
            for para in "\n".join(block).split("\n\n"):
                para = " ".join(para.split())
                if not para:
                    continue
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_after = Pt(8)
                add_runs(p, para, 12)

        elif line.startswith("|"):                               # pipe table
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip("|").split("|")]
                if not set("".join(cells)) <= set("-: "):
                    rows.append(cells)
                i += 1
            tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
            tbl.style = "Table Grid"
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r, cells in enumerate(rows):
                for c, txt in enumerate(cells):
                    cell = tbl.cell(r, c)
                    cell.paragraphs[0].text = ""
                    add_runs(cell.paragraphs[0], txt, 10.5, bold=(r == 0))
                    if c != 1:
                        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph()

        elif line.strip() == "---":                              # rule
            i += 1

        elif line.strip():                                       # paragraph
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            add_runs(p, line.strip(), 11,
                     color=MUTED if line.startswith("There are no") else INK)
            i += 1

        else:
            i += 1

    doc.save(OUT)
    print(f"wrote {OUT.name}  ({OUT.stat().st_size / 1e3:.0f} kB)")


if __name__ == "__main__":
    main()
